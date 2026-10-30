"""合同会社すまいる 尾﨑 葵 様 給与明細（令和8年7月分〜9月分）を生成する。

出力: kyuyo_meisai_R8_07-09.html / .pdf / .xlsx
前提:
  - 末締め・翌月25日支払、令和8年1月分から勤務、月給 270,000円（毎月同額）
  - 控除は雇用保険料と所得税のみ（健康保険・厚生年金・住民税なし）
  - 雇用保険料: 一般の事業・労働者負担。締日が令和8年4月1日以降の賃金は 5/1000、
    それ以前（1〜3月分）は令和7年度の 5.5/1000
  - 所得税: 令和8年分 月額表 甲欄、扶養親族等 0人。
    「電子計算機等を使用して源泉徴収税額を計算する方法」（電算機特例）で算出
  - 本年累計: 令和8年中の支払分（1月分〜当月分）の合計
"""
import subprocess
from decimal import Decimal, ROUND_HALF_UP, ROUND_CEILING
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).resolve().parent
COMPANY = "合同会社すまいる"
EMPLOYEE = "尾﨑　葵"
BASE_SALARY = 270_000
DEPENDENTS = 0
TARGET_MONTHS = [7, 8, 9]


def employment_insurance(month, wage):
    rate = Decimal("5.5") if month <= 3 else Decimal("5")  # 1/1000
    amount = Decimal(wage) * rate / 1000
    # 被保険者負担分の端数: 50銭以下切捨て、50銭超切上げ
    frac = amount - int(amount)
    return int(amount) + (1 if frac > Decimal("0.5") else 0)


def income_tax(taxable, dependents=DEPENDENTS):
    """令和8年分 月額表甲欄の電算機特例。"""
    s = Decimal(taxable)
    if s <= 158_333:
        kyuyo_kojo = Decimal(54_167)
    elif s <= 299_999:
        kyuyo_kojo = (s * Decimal("0.3")).to_integral_value(ROUND_CEILING) + 6_667
    elif s <= 549_999:
        kyuyo_kojo = (s * Decimal("0.2")).to_integral_value(ROUND_CEILING) + 36_667
    elif s <= 708_330:
        kyuyo_kojo = (s * Decimal("0.1")).to_integral_value(ROUND_CEILING) + 91_667
    else:
        kyuyo_kojo = Decimal(162_500)
    kiso_kojo = 48_334  # 月額 2,162,499円以下
    a = s - kyuyo_kojo - 31_667 * dependents - kiso_kojo
    a = max(Decimal(0), a)
    if a <= 162_500:
        tax = a * Decimal("0.05105")
    elif a <= 275_000:
        tax = a * Decimal("0.1021") - 8_296
    elif a <= 579_166:
        tax = a * Decimal("0.2042") - 36_374
    else:
        raise ValueError("対象外の金額")
    return int((tax / 10).quantize(Decimal(1), ROUND_HALF_UP) * 10)


def month_row(month):
    koyo = employment_insurance(month, BASE_SALARY)
    taxable = BASE_SALARY - koyo
    tax = income_tax(taxable)
    deduct = koyo + tax
    return {
        "month": month,
        "pay_date": f"令和8年{month + 1}月25日",
        "period": f"令和8年{month}月1日〜{month}月{[0,31,28,31,30,31,30,31,31,30,31,30,31][month]}日",
        "kihon": BASE_SALARY,
        "shikyu": BASE_SALARY,
        "koyo": koyo,
        "taxable": taxable,
        "tax": tax,
        "kojo": deduct,
        "net": BASE_SALARY - deduct,
    }


ROWS = {m: month_row(m) for m in range(1, 13)}


def cumulative(month):
    keys = ["shikyu", "koyo", "taxable", "tax", "kojo", "net"]
    return {k: sum(ROWS[m][k] for m in range(1, month + 1)) for k in keys}


def yen(n):
    return f"{n:,}"


# ---------- HTML / PDF ----------
CSS = """
@page { size: A4; margin: 14mm 16mm; }
* { box-sizing: border-box; }
body { font-family: 'IPAGothic', 'IPAPGothic', sans-serif; color: #222; font-size: 10.5pt; margin: 0; }
.slip { page-break-after: always; }
.slip:last-child { page-break-after: auto; }
h1 { text-align: center; font-size: 18pt; letter-spacing: .3em; margin: 0 0 4mm; }
.sub { text-align: center; margin-bottom: 6mm; font-size: 11pt; }
.head { display: flex; justify-content: space-between; align-items: flex-end; margin-bottom: 5mm; }
.name { font-size: 14pt; border-bottom: 1px solid #222; padding: 0 2mm 1mm; min-width: 70mm; }
.name small { font-size: 10pt; margin-left: 4mm; }
.company { text-align: right; line-height: 1.6; }
.company b { font-size: 12pt; }
table { width: 100%; border-collapse: collapse; margin-bottom: 4mm; table-layout: fixed; }
th, td { border: 1px solid #555; padding: 2mm 2.5mm; height: 9mm; }
th { background: #eef1f5; font-weight: normal; text-align: center; }
td.num { text-align: right; font-variant-numeric: tabular-nums; }
.side { width: 16mm; background: #dde3ec; writing-mode: vertical-rl; text-align: center; letter-spacing: .4em; }
.total th, .total td { font-weight: bold; }
.net th { background: #dde3ec; font-size: 12pt; }
.net td { font-size: 14pt; font-weight: bold; }
.caption { font-weight: bold; margin: 5mm 0 1.5mm; }
.note { font-size: 8.5pt; color: #555; line-height: 1.6; margin-top: 4mm; }
"""


def slip_html(r):
    c = cumulative(r["month"])
    return f"""
<section class="slip">
  <h1>給与明細書</h1>
  <div class="sub">令和8年{r['month']}月分（支給日：{r['pay_date']}）</div>
  <div class="head">
    <div class="name">{EMPLOYEE}<small>様</small></div>
    <div class="company"><b>{COMPANY}</b><br>対象期間：{r['period']}</div>
  </div>

  <table>
    <tr><td class="side" rowspan="2">支給</td>
      <th>基本給</th><th>時間外手当</th><th>通勤手当</th><th>その他</th><th>総支給額</th></tr>
    <tr><td class="num">{yen(r['kihon'])}</td><td class="num">0</td><td class="num">0</td><td class="num">0</td><td class="num"><b>{yen(r['shikyu'])}</b></td></tr>
  </table>

  <table>
    <tr><td class="side" rowspan="2">控除</td>
      <th>雇用保険料</th><th>所得税</th><th>住民税</th><th>その他</th><th>控除合計</th></tr>
    <tr><td class="num">{yen(r['koyo'])}</td><td class="num">{yen(r['tax'])}</td><td class="num">0</td><td class="num">0</td><td class="num"><b>{yen(r['kojo'])}</b></td></tr>
  </table>

  <table class="net">
    <tr><th style="width:40%">差引支給額</th><td class="num">{yen(r['net'])} 円</td></tr>
  </table>

  <div class="caption">本年累計（令和8年1月分〜{r['month']}月分・{r['month']}か月）</div>
  <table>
    <tr><th>総支給額</th><th>雇用保険料</th><th>課税支給額<br><small>（保険料控除後）</small></th><th>所得税</th><th>控除合計</th><th>差引支給額</th></tr>
    <tr><td class="num">{yen(c['shikyu'])}</td><td class="num">{yen(c['koyo'])}</td><td class="num">{yen(c['taxable'])}</td><td class="num">{yen(c['tax'])}</td><td class="num">{yen(c['kojo'])}</td><td class="num">{yen(c['net'])}</td></tr>
  </table>

  <div class="note">
    ※ 給与締日：毎月末日／支給日：翌月25日<br>
    ※ 雇用保険料：一般の事業 労働者負担 5/1000（令和8年度）。本年累計には令和7年度料率 5.5/1000 を適用した1〜3月分を含みます。<br>
    ※ 所得税：令和8年分 源泉徴収税額表 月額表 甲欄・扶養親族等 0人（電子計算機等による計算の特例により算出）
  </div>
</section>"""


def build_html():
    body = "\n".join(slip_html(ROWS[m]) for m in TARGET_MONTHS)
    return f"""<!doctype html>
<html lang="ja"><head><meta charset="utf-8"><title>給与明細書 令和8年7月〜9月分</title>
<style>{CSS}</style></head><body>{body}</body></html>"""


# ---------- Excel ----------
def build_xlsx(path):
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="555555")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="EEF1F5")
    side_fill = PatternFill("solid", fgColor="DDE3EC")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    font = "IPAゴシック"

    for m in TARGET_MONTHS:
        r, c = ROWS[m], cumulative(m)
        ws = wb.create_sheet(f"R8.{m}月分")
        ws.sheet_view.showGridLines = False
        for col, w in zip("ABCDEFG", [8, 15, 15, 15, 15, 15, 15]):
            ws.column_dimensions[col].width = w

        def put(cell, value, bold=False, size=11, fill=None, align=None, border=True, fmt=None):
            if value is not None:
                ws[cell] = value
            ws[cell].font = Font(name=font, bold=bold, size=size)
            if fill:
                ws[cell].fill = fill
            if align:
                ws[cell].alignment = align
            if border:
                ws[cell].border = box
            if fmt:
                ws[cell].number_format = fmt

        ws.merge_cells("A1:G1")
        put("A1", "給与明細書", bold=True, size=18, align=center, border=False)
        ws.merge_cells("A2:G2")
        put("A2", f"令和8年{m}月分（支給日：{r['pay_date']}）", align=center, border=False)
        ws.merge_cells("A4:C4")
        put("A4", f"{EMPLOYEE}　様", bold=True, size=14, border=False)
        ws.merge_cells("E4:G4")
        put("E4", COMPANY, bold=True, size=12, border=False, align=Alignment(horizontal="right"))
        ws.merge_cells("E5:G5")
        put("E5", f"対象期間：{r['period']}", border=False, align=Alignment(horizontal="right"))

        money = "#,##0"
        sections = [
            (7, "支給", ["基本給", "時間外手当", "通勤手当", "その他", "", "総支給額"],
             [r["kihon"], 0, 0, 0, None, "=SUM(B8:E8)"]),
            (10, "控除", ["雇用保険料", "所得税", "住民税", "その他", "", "控除合計"],
             [r["koyo"], r["tax"], 0, 0, None, "=SUM(B11:E11)"]),
        ]
        for row, label, heads, vals in sections:
            ws.merge_cells(f"A{row}:A{row + 1}")
            put(f"A{row}", label, fill=side_fill, align=center)
            put(f"A{row + 1}", None, fill=side_fill)
            for i, (h, v) in enumerate(zip(heads, vals)):
                col = "BCDEFG"[i]
                put(f"{col}{row}", h, fill=head_fill, align=center)
                put(f"{col}{row + 1}", v, fmt=money, bold=(i == 5))

        ws.merge_cells("A13:C13")
        put("A13", "差引支給額", bold=True, size=12, fill=side_fill, align=center)
        for col in "BC":
            put(f"{col}13", None, fill=side_fill)
        ws.merge_cells("D13:G13")
        put("D13", "=G8-G11", bold=True, size=14, fmt='#,##0"円"')
        for col in "EFG":
            put(f"{col}13", None)
        ws.row_dimensions[13].height = 26

        ws.merge_cells("A15:G15")
        put("A15", f"本年累計（令和8年1月分〜{m}月分・{m}か月）", bold=True, border=False)
        heads = ["", "総支給額", "雇用保険料", "課税支給額", "所得税", "控除合計", "差引支給額"]
        vals = [None, c["shikyu"], c["koyo"], c["taxable"], c["tax"], c["kojo"], c["net"]]
        for i, (h, v) in enumerate(zip(heads, vals)):
            col = "ABCDEFG"[i]
            if i == 0:
                ws.merge_cells("A16:A17")
                put("A16", "累計", fill=side_fill, align=center)
                put("A17", None, fill=side_fill)
                continue
            put(f"{col}16", h, fill=head_fill, align=center)
            put(f"{col}17", v, fmt=money)

        notes = [
            "※ 給与締日：毎月末日／支給日：翌月25日",
            "※ 雇用保険料：一般の事業 労働者負担 5/1000（令和8年度）。累計の1〜3月分は令和7年度料率 5.5/1000",
            "※ 所得税：令和8年分 月額表 甲欄・扶養親族等0人（電算機計算の特例により算出）",
        ]
        for i, n in enumerate(notes):
            ws.merge_cells(f"A{19 + i}:G{19 + i}")
            put(f"A{19 + i}", n, size=9, border=False)
        for rr in (8, 11, 17):
            ws.row_dimensions[rr].height = 22

        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = "A1:G21"

    wb.save(path)


def main():
    stem = OUT / "kyuyo_meisai_R8_07-09"
    html_path = stem.with_suffix(".html")
    html_path.write_text(build_html(), encoding="utf-8")
    subprocess.run(
        ["node", "-e", f"""
const {{ chromium }} = require('playwright');
(async () => {{
  const b = await chromium.launch();
  const p = await b.newPage();
  await p.goto('file://{html_path}');
  await p.pdf({{ path: '{stem.with_suffix(".pdf")}', format: 'A4', printBackground: true, preferCSSPageSize: true }});
  await b.close();
}})();
"""],
        check=True,
    )
    build_xlsx(stem.with_suffix(".xlsx"))

    print("月  雇用保険  課税対象   所得税  差引支給")
    for m in range(1, 13):
        r = ROWS[m]
        print(f"{m:>2}  {r['koyo']:>8,}  {r['taxable']:>9,}  {r['tax']:>7,}  {r['net']:>9,}")
    for m in TARGET_MONTHS:
        print(m, cumulative(m))


if __name__ == "__main__":
    main()
